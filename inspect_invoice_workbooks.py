import os
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.dirname(__file__))
files = [
    os.path.join(ROOT, '_cloud_template_Annapurna.xlsx'),
    os.path.join(ROOT, 'Generated_Invoices', 'Invoice_040-2026-27_Das_Metal.xlsx'),
    os.path.join(ROOT, 'Generated_Invoices', 'Invoice_036-2026-27_Tirupati_Udyog.xlsx'),
    os.path.join(ROOT, 'Generated_Invoices', 'Invoice_037-2026-27_Mirjamal.xlsx'),
    os.path.join(ROOT, 'Generated_Invoices', 'Invoice_038-2026-27_MSS.xlsx'),
    os.path.join(ROOT, 'Generated_Invoices', 'Invoice_039-2026-27_Mirjamal.xlsx'),
]

for path in files:
    print('=' * 80)
    print(f'Workbook: {path}')
    if not os.path.exists(path):
        print('  MISSING')
        continue

    wb = load_workbook(path, data_only=False)
    try:
        print('  sheets =', wb.sheetnames)
        for name in wb.sheetnames:
            sh = wb[name]
            print('  sheet:', name)
            print('    merged ranges:', [str(r) for r in sh.merged_cells.ranges])
            imgs = getattr(sh, '_images', [])
            print('    image count:', len(imgs))
            if imgs:
                print('    images:', [type(img).__name__ for img in imgs])
            for row in range(1, 26):
                cells = []
                for col in range(1, 9):
                    cell = sh.cell(row=row, column=col)
                    if cell.value not in (None, ''):
                        value = repr(cell.value)
                        if cell.data_type == 'f':
                            cells.append(f'{cell.coordinate}={value} [FORMULA]')
                        else:
                            cells.append(f'{cell.coordinate}={value}')
                if cells:
                    print('    ', row, '|', '; '.join(cells))
    finally:
        wb.close()
print('=' * 80)
